"""Evaluate the selected estimator, generate diagnostic plots, and build reports."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from sklearn.metrics import ConfusionMatrixDisplay, classification_report, confusion_matrix
from sklearn.pipeline import Pipeline


def evaluate_model(
    model: Pipeline,
    X_test: pd.DataFrame,
    y_test: pd.Series,
    output_dir: str | Path,
) -> dict[str, object]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    predictions = model.predict(X_test)
    report = classification_report(y_test, predictions, output_dict=True, zero_division=0)

    fig, ax = plt.subplots(figsize=(5, 4))
    ConfusionMatrixDisplay(
        confusion_matrix(y_test, predictions), display_labels=["No", "Yes"]
    ).plot(ax=ax, colorbar=False, cmap="Blues")
    ax.set_title("Confusion matrix")
    fig.tight_layout()
    fig.savefig(output_dir / "confusion_matrix.png", dpi=160)
    plt.close(fig)

    preprocessor = model.named_steps["preprocessor"]
    classifier = model.named_steps["classifier"]
    names = preprocessor.get_feature_names_out()
    if hasattr(classifier, "feature_importances_"):
        importance = classifier.feature_importances_
    elif hasattr(classifier, "coef_"):
        importance = np.abs(classifier.coef_[0])
    else:
        importance = np.zeros(len(names))
    top = pd.Series(importance, index=names).nlargest(20).sort_values()
    fig, ax = plt.subplots(figsize=(9, 7))
    top.plot.barh(ax=ax, color=sns.color_palette("viridis", len(top)))
    ax.set(title="Top feature importance", xlabel="Importance", ylabel="Feature")
    fig.tight_layout()
    fig.savefig(output_dir / "feature_importance.png", dpi=160)
    plt.close(fig)
    return report


def create_project_report(
    comparison: pd.DataFrame,
    classification: dict[str, object],
    report_path: str | Path,
    dataset_name: str,
    output_dir: str | Path | None = None,
) -> None:
    """Write model-comparison and test-set results with plots to a Word document."""
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Bank Term Deposit Prediction", level=0)
    document.add_paragraph(
        "This report summarizes a binary classification workflow for predicting whether "
        "a bank customer subscribes to a term deposit."
    )
    document.add_heading("Dataset", level=1)
    document.add_paragraph(
        f"Training data: {dataset_name}. The pipeline uses a stratified holdout test set (80/20) "
        "and fits preprocessing only on the training partition."
    )
    document.add_heading("Model Comparison", level=1)
    table = document.add_table(rows=1, cols=len(comparison.columns))
    table.style = "Light Shading Accent 1"
    for cell, column in zip(table.rows[0].cells, comparison.columns):
        cell.text = str(column)
    for _, row in comparison.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = f"{value:.4f}" if isinstance(value, (int, float)) else str(value)

    document.add_heading("Selected Model Test Results", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for cell, heading in zip(table.rows[0].cells, ("Class", "Precision", "Recall", "F1-score")):
        cell.text = heading
    for label, metrics in classification.items():
        if isinstance(metrics, dict) and "precision" in metrics:
            cells = table.add_row().cells
            values = (label, metrics["precision"], metrics["recall"], metrics["f1-score"])
            for cell, value in zip(cells, values):
                cell.text = f"{value:.4f}" if isinstance(value, (int, float)) else str(value)

    if output_dir is not None:
        out_path = Path(output_dir)
        document.add_heading("Visualizations", level=1)
        for img_name, title in [
            ("target_distribution.png", "Target Class Distribution"),
            ("age_vs_subscription.png", "Age vs Subscription"),
            ("job_vs_subscription.png", "Subscription Rate by Job"),
            ("confusion_matrix.png", "Confusion Matrix"),
            ("feature_importance.png", "Top Feature Importance"),
        ]:
            img_file = out_path / img_name
            if img_file.exists():
                document.add_heading(title, level=2)
                try:
                    document.add_picture(str(img_file), width=Inches(5.5))
                except Exception:
                    pass

    document.add_heading("Key Considerations", level=1)
    document.add_paragraph(
        "Note: the campaign-duration feature is included when present in the supplied data. "
        "It is only known after a call, so exclude it for pre-call predictions to avoid target leakage."
    )
    document.save(report_path)

